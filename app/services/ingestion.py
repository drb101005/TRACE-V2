import csv
import hashlib
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from fastapi import HTTPException
from app.core.errors import api_error

SUPPORTED = {".csv", ".xlsx"}
SALES_REQUIRED = ["invoice_number", "invoice_date", "customer_name", "total_amount"]
ALIASES = {
    "invoice_number": ["invoice number", "invoice no", "invoice no.", "invoice_number", "bill no", "document number"],
    "invoice_date": ["invoice date", "date", "invoice_date", "document date"],
    "customer_name": ["customer", "customer name", "buyer name", "party name", "customer_name"],
    "total_amount": ["total amount", "invoice amount", "gross amount", "total", "total_amount", "grand total"],
    "taxable_value": ["taxable value", "taxable amount", "taxable_value"],
    "tax_amount": ["tax amount", "gst amount", "tax", "tax_amount"],
}

def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()

def read_table(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as f:
                sample = f.read(8192)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.DictReader(f, dialect=dialect)
                if not reader.fieldnames:
                    raise api_error(422, "EMPTY_OR_INVALID_CSV", "CSV must contain a header row.")
                headers = [str(h).strip() for h in reader.fieldnames]
                rows = []
                for row in reader:
                    rows.append({str(k).strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k is not None})
                return headers, rows
        except UnicodeDecodeError:
            raise api_error(422, "CSV_ENCODING_INVALID", "CSV must be UTF-8 encoded. Save as UTF-8 CSV and retry.")
    if suffix == ".xlsx":
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            iterator = ws.iter_rows(values_only=True)
            first = next(iterator, None)
            if not first:
                raise api_error(422, "EMPTY_OR_INVALID_XLSX", "Workbook has no header row.")
            headers = [str(v).strip() if v is not None else "" for v in first]
            rows = []
            for values in iterator:
                if all(v is None or str(v).strip() == "" for v in values):
                    continue
                rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]})
            wb.close()
            return headers, rows
        except HTTPException:
            raise
        except Exception as exc:
            raise api_error(422, "XLSX_PARSE_FAILED", "Could not read this XLSX workbook. Check that it is a valid Excel file.")
    raise api_error(415, "UNSUPPORTED_FILE_TYPE", "Only CSV and XLSX files are supported in this phase.")

def normalized(s: Any) -> str:
    return " ".join(str(s or "").strip().lower().replace("_", " ").split())

def suggest_mapping(headers: list[str]) -> dict[str, str]:
    result = {}
    normalized_headers = {normalized(h): h for h in headers}
    for canonical, options in ALIASES.items():
        for option in options:
            if normalized(option) in normalized_headers:
                result[canonical] = normalized_headers[normalized(option)]
                break
    return result

def validate_rows(rows: list[dict[str, Any]], mapping: dict[str, str], dataset_type: str) -> dict[str, Any]:
    required = SALES_REQUIRED if dataset_type == "sales_register" else []
    missing = [field for field in required if not mapping.get(field)]
    errors = []
    seen = set()
    for idx, row in enumerate(rows, start=2):
        if len(errors) >= 500:
            break
        for field in missing:
            errors.append({"row_number": idx, "field": field, "code": "REQUIRED_COLUMN_UNMAPPED", "message": f"Required field '{field}' has no mapped source column."})
        if mapping.get("invoice_number"):
            invoice = str(row.get(mapping["invoice_number"], "") or "").strip()
            if not invoice:
                errors.append({"row_number": idx, "field": "invoice_number", "code": "REQUIRED_VALUE_MISSING", "message": "Invoice number is blank."})
            elif invoice.casefold() in seen:
                errors.append({"row_number": idx, "field": "invoice_number", "code": "DUPLICATE_IN_UPLOAD", "message": f"Invoice number '{invoice}' repeats in this upload."})
            else:
                seen.add(invoice.casefold())
        if mapping.get("total_amount"):
            raw = row.get(mapping["total_amount"])
            if raw is None or str(raw).strip() == "":
                errors.append({"row_number": idx, "field": "total_amount", "code": "AMOUNT_MISSING", "message": "Total amount is blank."})
            else:
                try:
                    amount = Decimal(str(raw).replace(",", "").replace("₹", "").strip())
                    if not amount.is_finite():
                        raise InvalidOperation
                except (InvalidOperation, ValueError):
                    errors.append({"row_number": idx, "field": "total_amount", "code": "INVALID_AMOUNT", "message": "Amount is not a valid decimal number."})
        if mapping.get("invoice_date"):
            raw_date = row.get(mapping["invoice_date"])
            if raw_date is None or str(raw_date).strip() == "":
                errors.append({"row_number": idx, "field": "invoice_date", "code": "DATE_MISSING", "message": "Invoice date is blank."})
    return {"valid": not errors and len(rows) > 0, "row_count": len(rows), "error_count": len(errors), "errors": errors,
            "warnings": ([{"code": "EMPTY_DATASET", "message": "No data rows were found."}] if not rows else []),
            "required_fields": required, "unmapped_required_fields": missing}

def canonicalize(row: dict[str, Any], mapping: dict[str, str]) -> dict[str, Any]:
    result = {"source": row}
    for canonical, source_col in mapping.items():
        value = row.get(source_col)
        if canonical in {"total_amount", "taxable_value", "tax_amount"} and value not in (None, ""):
            try:
                result[canonical] = format(Decimal(str(value).replace(",", "").replace("₹", "").strip()), "f")
            except (InvalidOperation, ValueError):
                result[canonical] = str(value)
        elif value is not None:
            result[canonical] = value.isoformat() if hasattr(value, "isoformat") else str(value).strip()
    return result
